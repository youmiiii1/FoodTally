import pandas as pd
import openpyxl
import tempfile
import os
from aiogram import Router, F
from aiogram.types import FSInputFile
from keyboards.meal_report import meal_report_menu
from keyboards.main_menu import main_menu_keyboard
from keyboards.photo_analyze import photo_analyze_keyboard
from texts.main_menu import main_menu_text
from texts.meal_report import (photo_analyze_wait, waiting_dish_name, wrong_dish_name, wrong_fat, wrong_carbs, fat_not_int, dish_name_not_text,
                               waiting_calories, waiting_protein, waiting_fat, waiting_carbs, wrong_protein, wrong_calories, carbs_not_int,
                               waiting_balance, looking_records, records_found, meal_report_menu_text, calories_not_int, protein_not_int, write_report_state_close)
from states.meal_report.write_info_state import WriteInfoStates
from states.photo_analyze.photo_take_states import PhotoTaker
from database.meal_report import new_report, show_all_reports, show_menu_reports
from openpyxl.styles import Alignment, Border, Side

router = Router()

@router.callback_query(F.data == 'meal_report')
async def meal_report_handler(call, db_pool):
    telegram_id = int(call.from_user.id)
    reports = await show_menu_reports(db_pool, telegram_id)

    await call.message.edit_text(f"{meal_report_menu_text}{reports}", reply_markup=meal_report_menu)
    await call.answer()

@router.callback_query(F.data == 'photo_report')
async def photo_analyze_handler(call, state):
    await state.set_state(PhotoTaker.waiting_for_photo)
    await state.update_data(mode='record')

    msg = await call.message.edit_text(photo_analyze_wait, reply_markup=photo_analyze_keyboard)
    await state.update_data(msg_id=msg.message_id)
    await state.update_data(chat_id=msg.chat.id)

    await call.answer()

@router.callback_query(F.data == 'to_menu')
async def to_main_menu_handler(call):
    await call.message.edit_text(main_menu_text, reply_markup=main_menu_keyboard)
    await call.answer()

@router.callback_query(F.data == 'get_records')
async def get_records_handler(call, db_pool):
    await call.message.answer(looking_records)
    telegram_id = int(call.from_user.id)
    reports = await show_all_reports(db_pool, telegram_id)

    columns = ["Dish Name", "Calories", "Protein_g", "Fat_g", "Carbs_g", "Balance Assessment", "Meal Time", "Meal Date", "Products"]
    df = pd.DataFrame(reports, columns=columns)

    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as temp:
        df.to_excel(temp.name, index=False, sheet_name='FoodTally Reports')

    wb = openpyxl.load_workbook(temp.name)
    ws = wb.active

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    # Formating table
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal='center',
                vertical='top',
                wrap_text=True
            )
            cell.border = border

        # Scale column with comments
    ws.column_dimensions['F'].width = 60

    # Set height of strokes (so text replace)
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 40

    # Scale and header pin
    ws.sheet_view.zoomScale = 130
    ws.freeze_panes = 'A2'

    wb.save(temp.name)
    path = temp.name

    await call.message.answer_document(FSInputFile(path, filename='FoodTally_Reports.xlsx'))
    os.remove(path)
    await call.message.answer(records_found)
    await call.answer()

@router.callback_query(F.data == 'write_report')
async def write_report_handler(call, state):
    await state.set_state(WriteInfoStates.waiting_for_dish_name)
    await call.message.answer(waiting_dish_name)
    await call.answer()

@router.message(WriteInfoStates.waiting_for_dish_name)
async def process_dish_name(message, state):
    try:
        dish_name = message.text.strip()

        if not dish_name or dish_name.isdigit():
            await message.answer(wrong_dish_name)
            return

        else:
            await state.update_data(dish_name=dish_name)
            await state.set_state(WriteInfoStates.waiting_for_calories_estimated)
            await message.answer(waiting_calories)

    except ValueError:
        await message.answer(dish_name_not_text)

@router.message(WriteInfoStates.waiting_for_calories_estimated)
async def process_calories_estimated(message, state):
    try:
        calories_estimated = int(message.text.strip())

        if calories_estimated <= 0 or calories_estimated >= 5000:
            await message.answer(wrong_calories)
            return

        else:
            await state.update_data(calories_estimated=calories_estimated)
            await state.set_state(WriteInfoStates.waiting_for_protein_g)
            await message.answer(waiting_protein)

    except ValueError:
        await message.answer(calories_not_int)

@router.message(WriteInfoStates.waiting_for_protein_g)
async def process_protein_g(message, state):
    try:
        protein_g = float(message.text.strip())

        if protein_g <= 0 or protein_g >= 500:
            await message.answer(wrong_protein)
            return

        else:
            await state.update_data(protein_g=protein_g)
            await state.set_state(WriteInfoStates.waiting_for_fat_g)
            await message.answer(waiting_fat)

    except ValueError:
        await message.answer(protein_not_int)

@router.message(WriteInfoStates.waiting_for_fat_g)
async def process_fat_g(message, state):
    try:
        fat_g = float(message.text.strip())

        if fat_g <= 0 or fat_g >= 500:
            await message.answer(wrong_fat)
            return

        else:
            await state.update_data(fat_g=fat_g)
            await state.set_state(WriteInfoStates.waiting_for_carbs_g)
            await message.answer(waiting_carbs)

    except ValueError:
        await message.answer(fat_not_int)

@router.message(WriteInfoStates.waiting_for_carbs_g)
async def process_carbs_g(message, state):
    try:
        carbs_g = float(message.text.strip())

        if carbs_g <= 0 or carbs_g >= 500:
            await message.answer(wrong_carbs)
            return

        else:
            await state.update_data(carbs_g=carbs_g)
            await state.set_state(WriteInfoStates.waiting_for_balance_assessment)
            await message.answer(waiting_balance)

    except ValueError:
        await message.answer(carbs_not_int)

@router.message(WriteInfoStates.waiting_for_balance_assessment)
async def process_balance_assessment(message, state, db_pool):
    try:
        balance_assessment = str(message.text.strip())

        if not balance_assessment:
            balance_assessment = None

        await state.update_data(balance_assessment=balance_assessment)

        data = await state.get_data()

        telegram_id = message.from_user.id
        dish_name = data['dish_name']
        calories_estimated = data['calories_estimated']
        protein_g = data['protein_g']
        fat_g = data['fat_g']
        carbs_g = data['carbs_g']
        balance_assessment = data['balance_assessment']

        await new_report(db_pool, telegram_id, dish_name, calories_estimated, protein_g, fat_g, carbs_g, balance_assessment)
        await state.clear()
        await message.answer(write_report_state_close)

        reports = await show_menu_reports(db_pool, telegram_id)
        await message.answer(f"{meal_report_menu_text}{reports}", reply_markup=meal_report_menu)

    except Exception as error:
        await message.answer(f"{error}")